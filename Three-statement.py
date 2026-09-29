import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "DCF & Valuation"
ws.views.sheetView[0].showGridLines = True

# Styling definitions (Institutional Standards)
font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=11, bold=True)
font_regular = Font(name="Calibri", size=11)
font_input = Font(name="Calibri", size=11, color="002060") # Wall Street Blue for Inputs

fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
fill_subtotal = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)
total_border = Border(
    top=Side(style="thin", color="000000"),
    bottom=Side(style="double", color="000000")
)

# Title Block
ws["A1"] = "TargetCo - 5-Year DCF & Valuation Model"
ws["A1"].font = font_title
ws["A2"] = "Amounts in USD Millions, unless otherwise noted"
ws["A2"].font = Font(name="Calibri", size=9, italic=True)

# Headers
headers = ["Line Item", "2025A", "2026E", "2027E", "2028E", "2029E", "2030E"]
for col_idx, text in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=col_idx, value=text)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left")

# Core Data & Formulas
rows_data = [
    # Label, Type, RowValues
    ("Revenue", "input", ["Revenue", 1000.0, 1100.0, 1210.0, 1331.0, 1464.1, 1610.5]),
    ("Revenue Growth %", "formula", [None, "-", "=C5/B5-1", "=D5/C5-1", "=E5/D5-1", "=F5/E5-1", "=G5/F5-1"]),
    ("COGS (60%)", "formula", [None, "=B5*0.60", "=C5*0.60", "=D5*0.60", "=E5*0.60", "=F5*0.60", "=G5*0.60"]),
    ("Gross Profit", "formula", [None, "=B5-B7", "=C5-C7", "=D5-D7", "=E5-E7", "=F5-F7", "=G5-G7"]),
    ("SG&A Expense (15%)", "formula", [None, "=B5*0.15", "=C5*0.15", "=D5*0.15", "=E5*0.15", "=F5*0.15", "=G5*0.15"]),
    ("EBITDA", "formula", [None, "=B8-B9", "=C8-C9", "=D8-D9", "=E8-E9", "=F8-F9", "=G8-G9"]),
    ("D&A Expense", "input", ["D&A", 40.0, 45.0, 50.0, 55.0, 60.0, 65.0]),
    ("Operating Income (EBIT)", "formula", [None, "=B10-B11", "=C10-C11", "=D10-D11", "=E10-E11", "=F10-F11", "=G10-G11"]),
    ("Tax Rate", "input", ["Tax Rate", 0.25, 0.25, 0.25, 0.25, 0.25, 0.25]),
    ("Taxes (NOPAT adjustment)", "formula", [None, "=B12*B13", "=C12*C13", "=D12*D13", "=E12*E13", "=F12*F13", "=G12*G13"]),
    ("NOPAT (EBIT * (1 - t))", "formula", [None, "=B12-B14", "=C12-C14", "=D12-D14", "=E12-E14", "=F12-F14", "=G12-G14"]),
    ("+ Depreciation & Amortization", "formula", [None, "=B11", "=C11", "=D11", "=E11", "=F11", "=G11"]),
    ("- Capital Expenditures (CapEx)", "input", ["CapEx", 50.0, 55.0, 60.0, 65.0, 70.0, 75.0]),
    ("- Change in Net Working Capital", "input", ["NWC", 15.0, 18.0, 20.0, 22.0, 24.0, 26.0]),
    ("Unlevered Free Cash Flow (UFCF)", "formula", [None, "=B15+B16-B17-B18", "=C15+C16-C17-C18", "=D15+D16-D17-D18", "=E15+E16-E17-E18", "=F15+F16-F17-F18", "=G15+G16-G17-G18"])
]

for row_idx, (label, row_type, values) in enumerate(rows_data, start=5):
    ws.cell(row=row_idx, column=1, value=label).font = font_bold if row_type == "formula" and "EBIT" in label or "Free" in label else font_regular
    for col_idx in range(2, 8):
        val = values[col_idx - 1]
        c = ws.cell(row=row_idx, column=col_idx, value=val)
        c.border = thin_border
        if "Growth" in label or "Tax Rate" in label:
            c.number_format = "0.0%"
            c.font = font_input if row_type == "input" else font_regular
        else:
            c.number_format = "$#,##0.0"
            c.font = font_input if row_type == "input" else font_regular

        if "Free Cash Flow" in label:
            c.fill = fill_subtotal
            c.font = font_bold
            c.border = total_border

# Valuation Parameter Block
params = [
    ("Discount Period (t)", 1, 2, 3, 4, 5),
    ("Discount Factor (WACC = 9.0%)", "=(1+0.09)^-C21", "=(1+0.09)^-D21", "=(1+0.09)^-E21", "=(1+0.09)^-F21", "=(1+0.09)^-G21"),
    ("Present Value of FCF", "=C19*C22", "=D19*D22", "=E19*E22", "=F19*F22", "=G19*G22")
]

ws.cell(row=21, column=1, value=params[0][0]).font = font_bold
for i, v in enumerate(params[0][1:], start=3):
    ws.cell(row=21, column=i, value=v).font = font_regular

ws.cell(row=22, column=1, value=params[1][0]).font = font_bold
for i, v in enumerate(params[1][1:], start=3):
    c = ws.cell(row=22, column=i, value=v)
    c.number_format = "0.000"
    c.font = font_regular

ws.cell(row=23, column=1, value=params[2][0]).font = font_bold
for i, v in enumerate(params[2][1:], start=3):
    c = ws.cell(row=23, column=i, value=v)
    c.number_format = "$#,##0.0"
    c.font = font_bold

# Enterprise to Equity Value Bridge
bridge = [
    ("Cumulative PV of 5-Yr FCF", "=SUM(C23:G23)"),
    ("Terminal Growth Rate (g)", 0.025),
    ("WACC", 0.09),
    ("Terminal Value (Gordon Growth)", "=(G19*(1+B26))/(B27-B26)"),
    ("PV of Terminal Value", "=B28*G22"),
    ("Enterprise Value (EV)", "=B25+B29"),
    ("Less: Total Debt", 350.0),
    ("Plus: Cash & Equivalents", 120.0),
    ("Implied Equity Value", "=B30-B31+B32"),
    ("Shares Outstanding (Diluted)", 15.0),
    ("Implied Target Share Price", "=B33/B34")
]

ws.cell(row=25, column=1, value="Valuation Summary Bridge").font = font_title
for r_offset, (label, val) in enumerate(bridge, start=25):
    ws.cell(row=r_offset, column=1, value=label).font = font_bold if "Implied" in label or "Enterprise" in label else font_regular
    c = ws.cell(row=r_offset, column=2, value=val)
    if "Rate" in label or "WACC" in label:
        c.number_format = "0.0%"
        c.font = font_input
    elif "Shares" in label:
        c.number_format = "#,##0.0"
        c.font = font_input
    elif "Price" in label:
        c.number_format = "$#,##0.00"
        c.font = font_bold
        c.border = total_border
    else:
        c.number_format = "$#,##0.0"
        c.font = font_bold if "Implied" in label or "Enterprise" in label else font_regular

for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

wb.save("TargetCo_Valuation_Model.xlsx")
print("Saved: TargetCo_Valuation_Model.xlsx")